import io
from openpyxl import load_workbook

from app.accounting.output import NBH_IMPORT_COLUMNS
from tests.test_accounting_workflow import (
    MEMBER_DATA,
    client_for,
    headers,
    records,
    start_job,
    token,
)


def test_human_corrections_endpoint_persists_user_edits():
    client, sheets, drive, _ = client_for()
    session_token = token(client)
    started = start_job(client, session_token, "MEMBER_RECEIPT")
    job = started.json()
    job_id = job["job_id"]

    auth = headers(session_token)

    # Initial row has raw bill head
    initial_rows = job["extracted_data"]["rows"]
    assert len(initial_rows) > 0

    # User explicitly edits row 0 in human review
    edited_rows = list(initial_rows)
    edited_rows[0] = dict(edited_rows[0])
    edited_rows[0]["Bill Head*"] = "Custom Maintenance Head"
    edited_rows[0]["Comments"] = "User edited comment"

    resp = client.post(
        f"/processing/jobs/{job_id}/corrections",
        headers=auth,
        json={"corrections": {"rows": edited_rows}},
    )
    assert resp.status_code == 200
    updated_job = resp.json()

    # User edit is authoritative and persisted
    updated_rows = updated_job["extracted_data"]["rows"]
    assert updated_rows[0]["Bill Head*"] == "Custom Maintenance Head"
    assert updated_rows[0]["Comments"] == "User edited comment"

    # Human corrections audit is recorded
    assert len(updated_job["human_corrections"]) > 0
    # Audit is per field and keyed by row_id (was one opaque "rows" entry).
    correction_fields = [c["field"] for c in updated_job["human_corrections"]]
    assert "Bill Head*" in correction_fields and "Comments" in correction_fields
    assert all(c["row_id"] == initial_rows[0]["_row_id"] for c in updated_job["human_corrections"])
    assert "Bill Head*" in updated_rows[0]["_edited_fields"]

    # Approve and verify Excel contains user-edited values
    approved = client.post(f"/processing/jobs/{job_id}/approve", headers=auth)
    assert approved.status_code == 200

    downloaded = client.get(f"/processing/jobs/{job_id}/download", headers=auth)
    assert downloaded.status_code == 200
    wb = load_workbook(io.BytesIO(downloaded.content))
    sheet = wb.active

    cols = [c.value for c in sheet[1]]
    assert cols == list(NBH_IMPORT_COLUMNS)
    assert sheet.cell(row=2, column=cols.index("Bill Head*") + 1).value == "Custom Maintenance Head"
    assert sheet.cell(row=2, column=cols.index("Comments") + 1).value == "User edited comment"


def test_unmapped_bill_head_remains_optional_without_blocking_approval():
    # Empty mapping table -> unmapped bill head
    client, _, drive, _ = client_for(sheet_records=records(include_mapping=False))
    session_token = token(client)
    started = start_job(client, session_token, "MEMBER_RECEIPT")
    job = started.json()

    assert job["overall_status"] == "NEEDS_REVIEW"
    assert job["mapping_status"] == "NEEDS_MAPPING"

    auth = headers(session_token)
    # Direct approval without resolving mapping must succeed
    approved = client.post(f"/processing/jobs/{job['job_id']}/approve", headers=auth)
    assert approved.status_code == 200
    assert approved.json()["overall_status"] == "COMPLETED"
