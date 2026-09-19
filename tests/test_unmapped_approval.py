import io
from openpyxl import load_workbook
import pytest

from app.accounting.output import NBH_IMPORT_COLUMNS
from tests.test_accounting_workflow import (
    MEMBER_DATA,
    client_for,
    headers,
    records,
    start_job,
    token,
)


def test_single_receipt_unmapped_bill_head_allows_approval_and_generates_excel():
    # Empty mapping table means bill head is unmapped
    client, _, drive, _ = client_for(sheet_records=records(include_mapping=False))
    session_token = token(client)
    started = start_job(client, session_token, "MEMBER_RECEIPT")
    job = started.json()

    # Verify job is in review and mapping is unconfirmed/needs mapping
    assert job["overall_status"] == "NEEDS_REVIEW"
    assert job["mapping_status"] == "NEEDS_MAPPING"
    assert len(job["mapping_result"]["missing"]) > 0

    # User does NOT resolve mapping, clicks approve directly
    auth = headers(session_token)
    approved = client.post(f"/processing/jobs/{job['job_id']}/approve", headers=auth)

    # Must succeed (no MAPPING_REQUIRED blocker)
    assert approved.status_code == 200
    completed_job = approved.json()
    assert completed_job["overall_status"] == "COMPLETED"
    assert completed_job["human_status"] == "APPROVED"

    # Verify Excel was generated and has exact 12 NBH columns with unmapped bill head preserved
    downloaded = client.get(f"/processing/jobs/{job['job_id']}/download", headers=auth)
    assert downloaded.status_code == 200
    wb = load_workbook(io.BytesIO(downloaded.content))
    sheet = wb.active
    columns = [cell.value for cell in sheet[1]]
    assert columns == list(NBH_IMPORT_COLUMNS)

    # First data row preserves original unmapped bill head ("Maintenance")
    first_row = [cell.value for cell in sheet[2]]
    bill_head_idx = list(NBH_IMPORT_COLUMNS).index("Bill Head*")
    assert first_row[bill_head_idx] == "Maintenance"
    assert [upload["folder_id"] for upload in drive.uploads] == ["member-in", "member-out"]
    assert len(drive.moves) == 1


def test_multi_row_unmapped_bill_heads_allows_approval_and_generates_excel():
    multi_row_data = {
        "rows": [
            {
                "Payment Type*": "Cash",
                "Society Bank Name/Bank code(Given to you by nobrokerhood)*": "Petty Cash",
                "Cheque/Ref No*": "REF-001",
                "Tower*": "Tower A",
                "Flat*": "101",
                "Bill Head*": "Property Tax",
                "Amount*": "1500.00",
                "Transaction Date*": "12-Sep-2026",
                "Comments": "Q1 Tax",
                "Cheque Date": "-",
                "Bank Name": "-",
                "Receipt No": "01",
            },
            {
                "Payment Type*": "Cash",
                "Society Bank Name/Bank code(Given to you by nobrokerhood)*": "Petty Cash",
                "Cheque/Ref No*": "REF-002",
                "Tower*": "Tower B",
                "Flat*": "202",
                "Bill Head*": "Water Charges",
                "Amount*": "850.00",
                "Transaction Date*": "12-Sep-2026",
                "Comments": "Sep Bill",
                "Cheque Date": "-",
                "Bank Name": "-",
                "Receipt No": "02",
            },
        ]
    }

    client, _, drive, _ = client_for(
        data={"MEMBER_RECEIPT": multi_row_data, "VENDOR_INVOICE": {}},
        sheet_records=records(include_mapping=False),
        # verification is per row_id: both rows must be covered
        verify_results=[{"overall_status": "PASSED", "rows": {"r0_001": "VERIFIED", "r0_002": "VERIFIED"}}],
    )
    session_token = token(client)
    started = start_job(client, session_token, "MEMBER_RECEIPT")
    job = started.json()

    assert job["overall_status"] == "NEEDS_REVIEW"
    assert job["mapping_status"] == "NEEDS_MAPPING"
    missing_sources = [m["source_value"] for m in job["mapping_result"]["missing"]]
    assert "Property Tax" in missing_sources
    assert "Water Charges" in missing_sources

    # Approve without confirming mappings
    auth = headers(session_token)
    approved = client.post(f"/processing/jobs/{job['job_id']}/approve", headers=auth)
    assert approved.status_code == 200
    assert approved.json()["overall_status"] == "COMPLETED"

    # Download Excel and check contents
    downloaded = client.get(f"/processing/jobs/{job['job_id']}/download", headers=auth)
    wb = load_workbook(io.BytesIO(downloaded.content))
    sheet = wb.active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0] == tuple(NBH_IMPORT_COLUMNS)
    assert len(rows) == 3  # Header + 2 rows

    bh_idx = list(NBH_IMPORT_COLUMNS).index("Bill Head*")
    assert rows[1][bh_idx] == "Property Tax"
    assert rows[2][bh_idx] == "Water Charges"


def test_genuine_blockers_still_prevent_approval():
    # 1. Verification not passed blocks approval
    unverified_client, _, _, _ = client_for(
        verify_results=[{"overall_status": "FAILED", "fields": []}], max_retries=0
    )
    token1 = token(unverified_client)
    job1 = start_job(unverified_client, token1).json()
    res1 = unverified_client.post(f"/processing/jobs/{job1['job_id']}/approve", headers=headers(token1))
    assert res1.status_code == 409
    # Blocked because the unverified row is an unresolved review item (not a
    # blanket "verification must be PASSED" rule a human could never satisfy).
    assert res1.json()["detail"]["code"] == "REVIEW_ITEMS_UNRESOLVED"
    assert any("ROW_NEEDS_REVIEW" in b for b in res1.json()["detail"]["blocking"])
    # ...and a human confirming the row resolves it.
    confirmed = unverified_client.post(f"/processing/jobs/{job1['job_id']}/rows/r0_001/confirm",
                                       headers=headers(token1), json={"reason": "checked against source"})
    assert confirmed.status_code == 200
    assert unverified_client.post(f"/processing/jobs/{job1['job_id']}/approve",
                                  headers=headers(token1)).status_code == 200

    # 2. Critical validation failure blocks approval
    invalid = MEMBER_DATA.copy()
    invalid["amount"] = None
    val_client, _, _, _ = client_for(
        data={"MEMBER_RECEIPT": invalid, "VENDOR_INVOICE": {}}
    )
    token2 = token(val_client)
    job2 = start_job(val_client, token2).json()
    assert job2["overall_status"] == "NEEDS_REVIEW"
    assert job2["validation_status"] == "BLOCKED"
    res2 = val_client.post(f"/processing/jobs/{job2['job_id']}/approve", headers=headers(token2))
    assert res2.status_code == 409
    assert any("AMOUNT_MISSING" in b for b in res2.json()["detail"]["blocking"])
    # Confirming the row does not bypass an invalid mandatory field.
    val_client.post(f"/processing/jobs/{job2['job_id']}/rows/r0_001/confirm", headers=headers(token2), json={})
    assert val_client.post(f"/processing/jobs/{job2['job_id']}/approve", headers=headers(token2)).status_code == 409
