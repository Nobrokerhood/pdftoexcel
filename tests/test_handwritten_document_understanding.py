import io
import pytest
from decimal import Decimal
from openpyxl import load_workbook

from app.accounting.purposes import MEMBER_RECEIPT
from app.accounting.reconciliation import AccountingReconciliationService
from app.accounting.schemas import MemberReceiptExtraction
from app.accounting.templates import MEMBER_RECEIPT_TEMPLATE, NBH_IMPORT_COLUMNS
from app.accounting.output import TemplateOutputGenerator
from app.agents.extractor import validate_extraction
from app.accounting.document_result import clean_cell as _clean_dash


def test_clean_dash_helper():
    assert _clean_dash(None) == "-"
    assert _clean_dash("") == "-"
    assert _clean_dash("null") == "-"
    assert _clean_dash("None") == "-"
    assert _clean_dash("N/A") == "-"
    assert _clean_dash("UNKNOWN") == "-"
    assert _clean_dash("  Office Expenses  ") == "Office Expenses"


def test_validate_extraction_normalizes_12_columns():
    raw_data = {
        "document_type": "PETTY_CASH_REGISTER",
        "summary": "Handwritten cash register",
        "rows": [
            {
                "Bill Head*": "Office",
                "Amount*": "900",
                "Transaction Date*": "04-07-25",
                "Cheque/Ref No*": "266",
                "Comments": "Firm Registration Stamp papers"
            },
            {
                "Bill Head*": "Salary",
                "Amount*": "10000",
                "Transaction Date*": "10-07-25",
                "Cheque/Ref No*": "284",
                "Comments": "Sujatha (H/K) Jun-25 Salary"
            }
        ]
    }
    validated = validate_extraction(MEMBER_RECEIPT, raw_data)
    rows = validated.get("rows", [])
    assert len(rows) == 2
    for r in rows:
        for col in NBH_IMPORT_COLUMNS:
            assert col in r
            assert r[col] is not None
            assert r[col] != ""


def test_reconciliation_service_math_and_discrepancy():
    recon_service = AccountingReconciliationService()
    extracted_data = {
        "rows": [
            {"Amount*": "900", "Cheque/Ref No*": "266"},
            {"Amount*": "140", "Cheque/Ref No*": "267"},
            {"Amount*": "700", "Cheque/Ref No*": "268"},
            {"Amount*": "10000", "Cheque/Ref No*": "284"},
            {"Amount*": "8400", "Cheque/Ref No*": "285"},
            {"Amount*": "9000", "Cheque/Ref No*": "286"},
            {"Amount*": "10000", "Cheque/Ref No*": "287"},
            {"Amount*": "9400", "Cheque/Ref No*": "288"},
            {"Amount*": "12000", "Cheque/Ref No*": "289"},
            {"Amount*": "10000", "Cheque/Ref No*": "290"},
        ],
        "balance_summary": {
            "opening_balance": "1714",
            "closing_balance": "10174",
            "total_expenditure": "98623",
            "total_receipts": "108800",
            "inflows": [
                {"date": "04-07-25", "amount": "10000", "ref_no": "014409"},
                {"date": "10-07-25", "amount": "68800", "ref_no": "024438"},
                {"date": "14-07-25", "amount": "30000", "ref_no": "014415"}
            ],
            "notes": ["Housekeeping salaries group sum: 68,800"]
        }
    }
    recon = recon_service.reconcile(extracted_data)
    checks = {c["check_id"]: c for c in recon["checks"]}
    assert recon["calculated_inflow_total"] == "108800"
    assert checks["RECEIPTS_TOTAL"]["status"] == "MATCHED"
    closing = checks["CLOSING_FROM_WRITTEN_TOTALS"]
    assert closing["source_value"] == "10174"
    assert closing["calculated_value"] == "10177"        # 108800 - 98623
    assert closing["difference"] == "-3"                  # source - calculated
    assert closing["status"] == "DISCREPANCY"


def test_exact_12_column_nbh_excel_output():
    generator = TemplateOutputGenerator()
    data = {
        "rows": [
            {
                "Payment Type*": "Cash",
                "Society Bank Name/Bank code(Given to you by nobrokerhood)*": "-",
                "Cheque/Ref No*": "266",
                "Tower No*": "-",
                "Flat No*": "-",
                "Bill Head*": "Office",
                "Amount*": "900",
                "Transaction Date*": "04-07-2025",
                "Comments": "Firm Registration Stamp papers",
                "Meter No": "-",
                "Cheque Issuer Bank": "-",
                "Cheque Date": "-"
            },
            {
                "Payment Type*": "Cash",
                "Society Bank Name/Bank code(Given to you by nobrokerhood)*": "-",
                "Cheque/Ref No*": "300",
                "Tower No*": "-",
                "Flat No*": "-",
                "Bill Head*": "Society Electrical",
                "Amount*": "2500",
                "Transaction Date*": "15-07-2025",
                "Comments": "Labour payment for common electrical wire (105)",
                "Meter No": "-",
                "Cheque Issuer Bank": "-",
                "Cheque Date": "-"
            }
        ]
    }
    filename, excel_bytes = generator.generate_xlsx(MEMBER_RECEIPT, MEMBER_RECEIPT_TEMPLATE, data, "test_job_123")
    wb = load_workbook(io.BytesIO(excel_bytes))
    sheet = wb.active

    # Check exact headers
    headers = [cell.value for cell in sheet[1]]
    assert headers == list(NBH_IMPORT_COLUMNS)
    assert len(headers) == 12

    # Check row counts and data
    assert sheet.max_row == 3
    assert sheet.cell(row=2, column=7).value == 900.0
    assert sheet.cell(row=3, column=7).value == 2500.0
    assert sheet.cell(row=3, column=3).value == "300"


def test_salary_group_vs_inflow_separation():
    # Demonstrates that salary group sum (68800) matches individual entries
    # and is distinct from payment entries.
    salary_rows = [
        {"name": "Sujatha", "amount": Decimal("10000")},
        {"name": "Saidamma", "amount": Decimal("8400")},
        {"name": "Lalitha", "amount": Decimal("9000")},
        {"name": "Krishnamma", "amount": Decimal("10000")},
        {"name": "Suvarna", "amount": Decimal("9400")},
        {"name": "Padma", "amount": Decimal("12000")},
        {"name": "Jhansi", "amount": Decimal("10000")},
    ]
    salary_sum = sum(s["amount"] for s in salary_rows)
    assert salary_sum == Decimal("68800")
    inflow_salary_amount = Decimal("68800")
    assert salary_sum == inflow_salary_amount


def test_handwritten_benchmark_full_regression():
    benchmark_rows = [
        {"Payment Type*": "Cash", "Society Bank Name/Bank code(Given to you by nobrokerhood)*": "-", "Cheque/Ref No*": "266", "Tower No*": "-", "Flat No*": "-", "Bill Head*": "Office", "Amount*": "900", "Transaction Date*": "04-07-2025", "Comments": "Firm Registration Stamp papers", "Meter No": "-", "Cheque Issuer Bank": "-", "Cheque Date": "-"},
        {"Payment Type*": "Cash", "Society Bank Name/Bank code(Given to you by nobrokerhood)*": "-", "Cheque/Ref No*": "267", "Tower No*": "-", "Flat No*": "-", "Bill Head*": "Office", "Amount*": "140", "Transaction Date*": "04-07-2025", "Comments": "Tea + Biscuits", "Meter No": "-", "Cheque Issuer Bank": "-", "Cheque Date": "-"},
        {"Payment Type*": "Cash", "Society Bank Name/Bank code(Given to you by nobrokerhood)*": "-", "Cheque/Ref No*": "268", "Tower No*": "-", "Flat No*": "-", "Bill Head*": "Office", "Amount*": "700", "Transaction Date*": "04-07-2025", "Comments": "Table cloth washing purpose", "Meter No": "-", "Cheque Issuer Bank": "-", "Cheque Date": "-"},
        {"Payment Type*": "Cash", "Society Bank Name/Bank code(Given to you by nobrokerhood)*": "-", "Cheque/Ref No*": "269", "Tower No*": "-", "Flat No*": "-", "Bill Head*": "Office", "Amount*": "130", "Transaction Date*": "04-07-2025", "Comments": "Tea + Biscuits (03-07-25)", "Meter No": "-", "Cheque Issuer Bank": "-", "Cheque Date": "-"},
        {"Payment Type*": "Cash", "Society Bank Name/Bank code(Given to you by nobrokerhood)*": "-", "Cheque/Ref No*": "270", "Tower No*": "-", "Flat No*": "-", "Bill Head*": "Shed (Office)", "Amount*": "1800", "Transaction Date*": "06-07-2025", "Comments": "Labour payment for cement work for shed", "Meter No": "-", "Cheque Issuer Bank": "-", "Cheque Date": "-"},
        {"Payment Type*": "Cash", "Society Bank Name/Bank code(Given to you by nobrokerhood)*": "-", "Cheque/Ref No*": "271", "Tower No*": "-", "Flat No*": "-", "Bill Head*": "Office", "Amount*": "300", "Transaction Date*": "07-07-2025", "Comments": "Water bottles 2 cans (office use)", "Meter No": "-", "Cheque Issuer Bank": "-", "Cheque Date": "-"},
        {"Payment Type*": "Cash", "Society Bank Name/Bank code(Given to you by nobrokerhood)*": "-", "Cheque/Ref No*": "274", "Tower No*": "-", "Flat No*": "-", "Bill Head*": "Garden", "Amount*": "200", "Transaction Date*": "08-07-2025", "Comments": "Grass cutting machine petrol", "Meter No": "-", "Cheque Issuer Bank": "-", "Cheque Date": "-"},
        {"Payment Type*": "Cash", "Society Bank Name/Bank code(Given to you by nobrokerhood)*": "-", "Cheque/Ref No*": "284", "Tower No*": "-", "Flat No*": "-", "Bill Head*": "Salary", "Amount*": "10000", "Transaction Date*": "10-07-2025", "Comments": "Sujatha (H/K) Jun-25 Salary", "Meter No": "-", "Cheque Issuer Bank": "-", "Cheque Date": "-"},
        {"Payment Type*": "Cash", "Society Bank Name/Bank code(Given to you by nobrokerhood)*": "-", "Cheque/Ref No*": "285", "Tower No*": "-", "Flat No*": "-", "Bill Head*": "Salary", "Amount*": "8400", "Transaction Date*": "10-07-2025", "Comments": "Saidamma (H/K) Jun-25 Salary", "Meter No": "-", "Cheque Issuer Bank": "-", "Cheque Date": "-"},
        {"Payment Type*": "Cash", "Society Bank Name/Bank code(Given to you by nobrokerhood)*": "-", "Cheque/Ref No*": "286", "Tower No*": "-", "Flat No*": "-", "Bill Head*": "Salary", "Amount*": "9000", "Transaction Date*": "10-07-2025", "Comments": "Lalitha (H/K) Jun-25 Salary", "Meter No": "-", "Cheque Issuer Bank": "-", "Cheque Date": "-"},
        {"Payment Type*": "Cash", "Society Bank Name/Bank code(Given to you by nobrokerhood)*": "-", "Cheque/Ref No*": "287", "Tower No*": "-", "Flat No*": "-", "Bill Head*": "Salary", "Amount*": "10000", "Transaction Date*": "10-07-2025", "Comments": "Krishnamma (H/K) Jun-25 Salary", "Meter No": "-", "Cheque Issuer Bank": "-", "Cheque Date": "-"},
        {"Payment Type*": "Cash", "Society Bank Name/Bank code(Given to you by nobrokerhood)*": "-", "Cheque/Ref No*": "288", "Tower No*": "-", "Flat No*": "-", "Bill Head*": "Salary", "Amount*": "9400", "Transaction Date*": "10-07-2025", "Comments": "Suvarna (H/K) Jun-25 Salary", "Meter No": "-", "Cheque Issuer Bank": "-", "Cheque Date": "-"},
        {"Payment Type*": "Cash", "Society Bank Name/Bank code(Given to you by nobrokerhood)*": "-", "Cheque/Ref No*": "289", "Tower No*": "-", "Flat No*": "-", "Bill Head*": "Salary", "Amount*": "12000", "Transaction Date*": "10-07-2025", "Comments": "Padma (H/K) Jun-25 Salary", "Meter No": "-", "Cheque Issuer Bank": "-", "Cheque Date": "-"},
        {"Payment Type*": "Cash", "Society Bank Name/Bank code(Given to you by nobrokerhood)*": "-", "Cheque/Ref No*": "290", "Tower No*": "-", "Flat No*": "-", "Bill Head*": "Salary", "Amount*": "10000", "Transaction Date*": "10-07-2025", "Comments": "Jhansi (H/K) Jun-25 Salary", "Meter No": "-", "Cheque Issuer Bank": "-", "Cheque Date": "-"},
        {"Payment Type*": "Cash", "Society Bank Name/Bank code(Given to you by nobrokerhood)*": "-", "Cheque/Ref No*": "291", "Tower No*": "-", "Flat No*": "-", "Bill Head*": "Society Electrical", "Amount*": "880", "Transaction Date*": "10-07-2025", "Comments": "common electrical wire Transport charges", "Meter No": "-", "Cheque Issuer Bank": "-", "Cheque Date": "-"},
        {"Payment Type*": "Cash", "Society Bank Name/Bank code(Given to you by nobrokerhood)*": "-", "Cheque/Ref No*": "292", "Tower No*": "-", "Flat No*": "-", "Bill Head*": "Garden", "Amount*": "200", "Transaction Date*": "10-07-2025", "Comments": "Grass cutting machine Petrol", "Meter No": "-", "Cheque Issuer Bank": "-", "Cheque Date": "-"},
        {"Payment Type*": "Cash", "Society Bank Name/Bank code(Given to you by nobrokerhood)*": "-", "Cheque/Ref No*": "296", "Tower No*": "-", "Flat No*": "-", "Bill Head*": "Water", "Amount*": "1090", "Transaction Date*": "15-07-2025", "Comments": "Suresh water supply drinking water 10 cans", "Meter No": "-", "Cheque Issuer Bank": "-", "Cheque Date": "-"},
        {"Payment Type*": "Cash", "Society Bank Name/Bank code(Given to you by nobrokerhood)*": "-", "Cheque/Ref No*": "297", "Tower No*": "-", "Flat No*": "-", "Bill Head*": "Office (Shed)", "Amount*": "2250", "Transaction Date*": "15-07-2025", "Comments": "Krishna Reddy Iron Stand Rent for (Shed)", "Meter No": "-", "Cheque Issuer Bank": "-", "Cheque Date": "-"},
        {"Payment Type*": "Cash", "Society Bank Name/Bank code(Given to you by nobrokerhood)*": "-", "Cheque/Ref No*": "298", "Tower No*": "-", "Flat No*": "-", "Bill Head*": "Office (Shed)", "Amount*": "5000", "Transaction Date*": "15-07-2025", "Comments": "Md. Nasar Table Advance for (shed)", "Meter No": "-", "Cheque Issuer Bank": "-", "Cheque Date": "-"},
        {"Payment Type*": "Cash", "Society Bank Name/Bank code(Given to you by nobrokerhood)*": "-", "Cheque/Ref No*": "299", "Tower No*": "-", "Flat No*": "-", "Bill Head*": "Miscellaneous", "Amount*": "2000", "Transaction Date*": "15-07-2025", "Comments": "Iron water work miscellaneous (Jun-25)", "Meter No": "-", "Cheque Issuer Bank": "-", "Cheque Date": "-"},
        {"Payment Type*": "Cash", "Society Bank Name/Bank code(Given to you by nobrokerhood)*": "-", "Cheque/Ref No*": "300", "Tower No*": "-", "Flat No*": "-", "Bill Head*": "Society Electrical", "Amount*": "2500", "Transaction Date*": "15-07-2025", "Comments": "Labour payment for common electrical wire (105)", "Meter No": "-", "Cheque Issuer Bank": "-", "Cheque Date": "-"},
        {"Payment Type*": "Cash", "Society Bank Name/Bank code(Given to you by nobrokerhood)*": "-", "Cheque/Ref No*": "301", "Tower No*": "-", "Flat No*": "-", "Bill Head*": "Shed (Office)", "Amount*": "2400", "Transaction Date*": "15-07-2025", "Comments": "Iron stand Rent purpose for shed 8 days", "Meter No": "-", "Cheque Issuer Bank": "-", "Cheque Date": "-"},
        {"Payment Type*": "Cash", "Society Bank Name/Bank code(Given to you by nobrokerhood)*": "-", "Cheque/Ref No*": "302", "Tower No*": "-", "Flat No*": "-", "Bill Head*": "Society", "Amount*": "1789", "Transaction Date*": "15-07-2025", "Comments": "Kiranam & General Stores (water bottles & sunblock)", "Meter No": "-", "Cheque Issuer Bank": "-", "Cheque Date": "-"},
        {"Payment Type*": "Cash", "Society Bank Name/Bank code(Given to you by nobrokerhood)*": "-", "Cheque/Ref No*": "303", "Tower No*": "-", "Flat No*": "-", "Bill Head*": "Miscellaneous", "Amount*": "1000", "Transaction Date*": "15-07-2025", "Comments": "Union Bank A/c Transfer purpose (Attender)", "Meter No": "-", "Cheque Issuer Bank": "-", "Cheque Date": "-"},
        {"Payment Type*": "Cash", "Society Bank Name/Bank code(Given to you by nobrokerhood)*": "-", "Cheque/Ref No*": "304", "Tower No*": "-", "Flat No*": "-", "Bill Head*": "Salary Advance", "Amount*": "3000", "Transaction Date*": "15-07-2025", "Comments": "Jhansi (H/K) Salary Advance (July-25)", "Meter No": "-", "Cheque Issuer Bank": "-", "Cheque Date": "-"},
        {"Payment Type*": "Cash", "Society Bank Name/Bank code(Given to you by nobrokerhood)*": "-", "Cheque/Ref No*": "305", "Tower No*": "-", "Flat No*": "-", "Bill Head*": "Society", "Amount*": "200", "Transaction Date*": "15-07-2025", "Comments": "Mahender mosquito fogging", "Meter No": "-", "Cheque Issuer Bank": "-", "Cheque Date": "-"},
        {"Payment Type*": "Cash", "Society Bank Name/Bank code(Given to you by nobrokerhood)*": "-", "Cheque/Ref No*": "306", "Tower No*": "-", "Flat No*": "-", "Bill Head*": "Garden", "Amount*": "200", "Transaction Date*": "18-07-2025", "Comments": "Grass cutting machine Petrol", "Meter No": "-", "Cheque Issuer Bank": "-", "Cheque Date": "-"},
        {"Payment Type*": "Cash", "Society Bank Name/Bank code(Given to you by nobrokerhood)*": "-", "Cheque/Ref No*": "307", "Tower No*": "-", "Flat No*": "-", "Bill Head*": "Shed (Office)", "Amount*": "750", "Transaction Date*": "19-07-2025", "Comments": "Krishna Reddy Iron Stand Rent 1 day", "Meter No": "-", "Cheque Issuer Bank": "-", "Cheque Date": "-"},
        {"Payment Type*": "Cash", "Society Bank Name/Bank code(Given to you by nobrokerhood)*": "-", "Cheque/Ref No*": "308", "Tower No*": "-", "Flat No*": "-", "Bill Head*": "Shed (Office)", "Amount*": "200", "Transaction Date*": "19-07-2025", "Comments": "Nuts & Bolts for fitting purpose (shed)", "Meter No": "-", "Cheque Issuer Bank": "-", "Cheque Date": "-"},
    ]

    assert len(benchmark_rows) == 29
    assert benchmark_rows[20]["Amount*"] == "2500"
    assert benchmark_rows[20]["Comments"] == "Labour payment for common electrical wire (105)"
    assert benchmark_rows[22]["Amount*"] == "1789"
    assert benchmark_rows[22]["Comments"] == "Kiranam & General Stores (water bottles & sunblock)"
    assert benchmark_rows[27]["Amount*"] == "750"
    assert benchmark_rows[27]["Comments"] == "Krishna Reddy Iron Stand Rent 1 day"
    assert benchmark_rows[13]["Cheque/Ref No*"] == "290"

    inflows = [
        {"date": "04-07-25", "amount": "10000", "ref_no": "014409"},
        {"date": "10-07-25", "amount": "68800", "ref_no": "024438"},
        {"date": "14-07-25", "amount": "30000", "ref_no": "014415"},
    ]
    assert len(inflows) == 3
    inflow_total = sum(float(x["amount"]) for x in inflows)
    assert inflow_total == 108800.0

    extracted_data = {
        "rows": benchmark_rows,
        "balance_summary": {
            "opening_balance": "1714",
            "closing_balance": "10174",
            "total_expenditure": "98623",
            "total_receipts": "108800",
            "inflows": inflows,
        }
    }

    # Golden reconciliation: source-written vs calculated, never adjusted.
    extracted_data["balance_summary"]["opening_balance"] = "-1714"   # written as a deficit
    before = [dict(r) for r in benchmark_rows]
    recon = AccountingReconciliationService().reconcile(extracted_data)
    checks = {c["check_id"]: c for c in recon["checks"]}
    assert recon["calculated_transaction_total"] == "96429"
    assert recon["calculated_inflow_total"] == "108800"
    assert checks["EXPENDITURE_TOTAL"]["source_value"] == "98623"
    assert checks["EXPENDITURE_TOTAL"]["status"] == "DISCREPANCY"
    assert recon["derived_opening_adjustment"] == "2194"
    closing = checks["CLOSING_FROM_WRITTEN_TOTALS"]
    assert (closing["source_value"], closing["calculated_value"], closing["difference"]) == ("10174", "10177", "-3")
    opening = checks["OPENING_ADJUSTMENT"]
    assert (opening["source_value"], opening["calculated_value"], opening["difference"]) == ("1714", "2194", "-480")
    assert opening["status"] == closing["status"] == "DISCREPANCY"
    assert recon["overall_status"] == "DISCREPANCY"
    assert benchmark_rows == before, "reconciliation must never alter source values"

    # Both discrepancies surface as non-blocking review warnings.
    from app.accounting.validation import AccountingValidationService
    extracted_data["reconciliation"] = recon
    for row in extracted_data["rows"]:
        row["_status"] = "ACCEPTED"
    issues = AccountingValidationService().validate("PETTY_CASH_REGISTER", extracted_data).issues
    warned = [i for i in issues if i.code == "RECONCILIATION_DISCREPANCY"]
    assert any("-3" in i.message for i in warned) and any("-480" in i.message for i in warned)
    assert all(i.severity == "WARNING" for i in warned)

    generator = TemplateOutputGenerator()
    _, excel_bytes = generator.generate_xlsx(MEMBER_RECEIPT, MEMBER_RECEIPT_TEMPLATE, extracted_data, "job_bench")
    wb = load_workbook(io.BytesIO(excel_bytes))
    ws = wb.active
    assert ws.max_row == 30
    assert ws.cell(row=22, column=7).value == 2500.0
    assert ws.cell(row=24, column=7).value == 1789.0
    assert ws.cell(row=29, column=7).value == 750.0
    assert ws.cell(row=15, column=3).value == "290"

