"""Approval guards, exact XLSX output, bounded field-level repair.

Rewritten for the row-id review model: approval is blocked by unresolved review
items (not by a blanket 'verification must be PASSED' that no human could
satisfy); amounts/dates go through the strict parsers.
"""

import io
from datetime import date, datetime

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from app.accounting.money import parse_amount
from app.accounting.output import OutputGenerationError, nbh_cells
from app.accounting.schemas import MemberReceiptExtraction, VendorInvoiceExtraction
from app.core.errors import ExternalServiceUnavailableError
from app.google.sheets_service import GoogleSheetsService
from tests.test_accounting_workflow import (
    MEMBER_DATA,
    VENDOR_DATA,
    StaticRepairProvider,
    client_for,
    headers,
    settings,
    start_job,
    token,
)


class UnavailableVerifier:
    def verify(self, source_bytes, purpose, template, extracted_data):
        raise ExternalServiceUnavailableError("verification service down")


class FailingExtraction:
    def extract(self, source_bytes, purpose, template):
        raise ExternalServiceUnavailableError("extraction service down")


def test_verification_outage_leaves_rows_unverified_and_blocks_approval():
    client, app, drive, _ = client_for()
    app.state.verification_agent.provider = UnavailableVerifier()
    session_token = token(client)

    job = start_job(client, session_token).json()
    assert job["overall_status"] == "NEEDS_REVIEW"
    assert job["verification_status"] == "NEEDS_REVIEW"          # never PASSED
    assert job["extracted_data"]["rows"][0]["_status"] == "UNVERIFIED"

    approved = client.post(f"/processing/jobs/{job['job_id']}/approve", headers=headers(session_token))
    assert approved.status_code == 409
    assert approved.json()["detail"]["code"] == "REVIEW_ITEMS_UNRESOLVED"
    assert [u["folder_id"] for u in drive.uploads] == ["member-in"]
    assert drive.moves == []


def test_failed_job_cannot_be_approved_but_can_be_rejected():
    client, app, drive, _ = client_for()
    app.state.extraction_agent.provider = FailingExtraction()
    session_token = token(client)
    job = start_job(client, session_token).json()
    assert job["overall_status"] == "FAILED"
    stages = {p["stage"]: p["status"] for p in job["progress"]}
    assert stages["OCR_AND_EXTRACTION"] == "FAILED"
    assert "VERIFICATION" not in stages                              # never claimed

    approved = client.post(f"/processing/jobs/{job['job_id']}/approve", headers=headers(session_token))
    assert approved.status_code == 409
    assert approved.json()["detail"]["code"] == "JOB_NOT_AWAITING_REVIEW"
    rejected = client.post(f"/processing/jobs/{job['job_id']}/reject", headers=headers(session_token))
    assert rejected.status_code == 200 and rejected.json()["overall_status"] == "REJECTED"
    assert drive.moves == []


def test_completed_job_cannot_be_approved_again_or_rejected():
    client, _, drive, _ = client_for()
    session_token = token(client)
    job = start_job(client, session_token).json()
    auth = headers(session_token)

    assert client.post(f"/processing/jobs/{job['job_id']}/approve", headers=auth).status_code == 200
    assert client.post(f"/processing/jobs/{job['job_id']}/approve", headers=auth).status_code == 409
    assert client.post(f"/processing/jobs/{job['job_id']}/reject", headers=auth).status_code == 409
    assert [u["folder_id"] for u in drive.uploads] == ["member-in", "member-out"]
    assert len(drive.moves) == 1


def download_book(client, session_token, purpose):
    job = start_job(client, session_token, purpose).json()
    auth = headers(session_token)
    resp = client.post(f"/processing/jobs/{job['job_id']}/approve", headers=auth)
    assert resp.status_code == 200, resp.text
    content = client.get(f"/processing/jobs/{job['job_id']}/download", headers=auth).content
    return load_workbook(io.BytesIO(content))


def test_member_receipt_xlsx_writes_numeric_amount_and_real_date():
    client, _, _, _ = client_for()
    wb = download_book(client, token(client), "MEMBER_RECEIPT")
    sheet = wb.active
    header = [c.value for c in sheet[1]]
    row = dict(zip(header, next(sheet.iter_rows(min_row=2))))
    assert row["Amount*"].data_type == "n" and row["Amount*"].value == 5000.0
    assert isinstance(row["Transaction Date*"].value, datetime)
    assert row["Transaction Date*"].value.date() == date(2026, 8, 25)
    assert row["Transaction Date*"].number_format == "DD-MM-YYYY"
    assert row["Cheque/Ref No*"].value == MEMBER_DATA["reference_number"]
    assert row["Meter No"].value == "-"
    assert {"Evidence", "Reconciliation", "Source Candidates", "Review Audit", "Summary"} <= set(wb.sheetnames)


def test_vendor_invoice_xlsx_is_12_nbh_columns_with_vendor_detail_sheet():
    client, _, _, _ = client_for(verify_results=[{"overall_status": "PASSED", "rows": {"r0_001": "VERIFIED"}}])
    wb = download_book(client, token(client), "VENDOR_INVOICE")
    sheet = wb.active
    header = [c.value for c in sheet[1]]
    assert len(header) == 12 and header[6] == "Amount*"
    row = dict(zip(header, next(sheet.iter_rows(min_row=2))))
    assert row["Amount*"].value == 10000.0
    assert row["Cheque/Ref No*"].value == VENDOR_DATA["bill_number"]
    assert row["Transaction Date*"].value.date() == date(2026, 8, 25)


@pytest.mark.parametrize("raw, expected", [
    ("Rs.900", 900.0), ("₹ 12,450.75", 12450.75), ("Rs. 1,789/-", 1789.0), ("750/-", 750.0),
    ("1,00,000", 100000.0), ("0", 0.0),
])
def test_export_amounts_are_parsed_strictly(raw, expected):
    row = {"Amount*": raw, "Transaction Date*": "04-07-2025"}
    cells = dict(zip(range(12), nbh_cells(row, "r")))
    assert cells[6] == ("amount", expected)


@pytest.mark.parametrize("raw", ["2500 (105)", "1.500,00", "18001-", "see attached", "-"])
def test_export_refuses_ambiguous_or_missing_amounts(raw):
    """Old writer: 'Rs.900'->0.9, '2500 (105)'->2500105, '1.500,00'->1.5. Now refused."""
    with pytest.raises(OutputGenerationError):
        nbh_cells({"Amount*": raw, "Transaction Date*": "04-07-2025"}, "r")


def test_export_dates_follow_one_policy():
    for raw in ("04-07-25", "04/07/2025", "2025-07-04", "4 Jul 2025"):
        kind, value = nbh_cells({"Amount*": "1", "Transaction Date*": raw}, "r")[7]
        assert (kind, value) == ("date", date(2025, 7, 4)), raw
    with pytest.raises(OutputGenerationError):
        nbh_cells({"Amount*": "1", "Transaction Date*": "next Tuesday"}, "r")
    assert nbh_cells({"Amount*": "1", "Transaction Date*": "04-07-25", "Cheque Date": "-"}, "r")[11] == ("text", "-")


def test_formula_like_source_text_is_stored_as_text():
    from app.accounting.output import TemplateOutputGenerator
    from app.accounting.templates import MEMBER_RECEIPT_TEMPLATE

    data = {"rows": [{"Amount*": "100", "Transaction Date*": "04-07-2025", "Comments": '=HYPERLINK("http://x","y")',
                      "Bill Head*": "+cmd|' /C calc'!A0", "Cheque/Ref No*": "-", "_row_id": "r1"}]}
    _, content = TemplateOutputGenerator().generate_xlsx("MEMBER_RECEIPT", MEMBER_RECEIPT_TEMPLATE, data, "j")
    ws = load_workbook(io.BytesIO(content)).active
    header = [c.value for c in ws[1]]
    for col in ("Comments", "Bill Head*"):
        cell = ws.cell(row=2, column=header.index(col) + 1)
        assert cell.data_type == "s", col
    assert ws.cell(row=2, column=header.index("Cheque/Ref No*") + 1).value == "-"


def test_zero_rows_never_produce_a_workbook():
    from app.accounting.output import TemplateOutputGenerator
    from app.accounting.templates import MEMBER_RECEIPT_TEMPLATE

    with pytest.raises(OutputGenerationError):
        TemplateOutputGenerator().generate_xlsx("MEMBER_RECEIPT", MEMBER_RECEIPT_TEMPLATE,
                                                {"rows": [], "amount": "5000", "reference_number": "X"}, "j")


class RecordingWorksheet:
    def __init__(self, rows):
        self.rows = rows
        self.batches = []

    def get_all_values(self):
        return self.rows

    def batch_update(self, data, value_input_option=None):
        self.batches.append((data, value_input_option))

    def update_cell(self, *args):
        raise AssertionError("per-cell writes exhaust the Sheets write quota")


def test_row_update_is_a_single_batched_raw_write():
    worksheet = RecordingWorksheet(
        [["Job ID", "Overall Status", "Completed At"], ["job-1", "PROCESSING", ""], ["job-2", "CREATED", ""]])
    service = GoogleSheetsService(settings())
    service._worksheet = lambda spreadsheet_id, worksheet_name="Sheet1": worksheet
    assert service.update_row_by_key("sheet", "Job ID", "job-2",
                                     {"Overall Status": "COMPLETED", "Completed At": "now", "Unknown": "x"}) is True
    assert len(worksheet.batches) == 1
    data, input_option = worksheet.batches[0]
    assert data == [{"range": "B3", "values": [["COMPLETED"]]}, {"range": "C3", "values": [["now"]]}]
    # RAW: user-controlled text is never evaluated as a Sheets formula.
    assert str(input_option.value if hasattr(input_option, "value") else input_option) == "RAW"


@pytest.mark.parametrize("raw, expected", [
    ("12,450.75", "12450.75"), ("₹ 12,450.75", "12450.75"), ("Rs. 1,00,000", "100000"),
    ("INR 500", "500"), ("5000", "5000"), (5000, "5000"), ("", None), (None, None),
])
def test_extraction_schema_accepts_printed_amount_formats(raw, expected):
    assert MemberReceiptExtraction(amount=raw).model_dump(mode="json")["amount"] == expected


@pytest.mark.parametrize("raw", ["twelve thousand", "12.450,75", "Rs 12 abc", "PRS 100"])
def test_extraction_schema_rejects_ambiguous_amounts(raw):
    with pytest.raises(ValidationError):
        MemberReceiptExtraction(amount=raw)
    assert not parse_amount(raw).found


def test_vendor_schema_normalizes_tax_and_expense_amounts():
    data = VendorInvoiceExtraction(cgst_amount="1,125.00",
                                   expenses=[{"expense_amount": "₹10,000"}]).model_dump(mode="json")
    assert data["cgst_amount"] == "1125.00"
    assert data["expenses"][0]["expense_amount"] == "10000"


LIVE_MALFORMED_VERIFICATION = {
    "overall_status": "PASSED",
    "fields": [{"field": "amount", "extracted_value": "5000", "verified_value": "5000",
                "status": "PASSED", "confidence": "HIGH", "evidence": "Amount Paid"}],
}


def test_malformed_verification_is_never_passed_and_never_500():
    client, app, drive, _ = client_for(verify_results=[LIVE_MALFORMED_VERIFICATION])
    repair = app.state.repair_agent.provider
    session_token = token(client)
    response = start_job(client, session_token)
    assert response.status_code == 200
    job = response.json()
    assert job["overall_status"] == "NEEDS_REVIEW"
    assert job["verification_status"] == "NEEDS_REVIEW"
    assert repair.calls == 0
    assert client.post(f"/processing/jobs/{job['job_id']}/approve", headers=headers(session_token)).status_code == 409
    assert len(drive.uploads) == 1


def test_repair_only_fills_the_requested_field_and_normalizes_it():
    missing = dict(MEMBER_DATA, amount=None)
    client, app, _, _ = client_for(data={"MEMBER_RECEIPT": missing, "VENDOR_INVOICE": VENDOR_DATA},
                                   verify_results=[{"overall_status": "FAILED", "fields": []}])
    app.state.repair_agent.provider = StaticRepairProvider({"Amount*": "₹ 5,000.00", "Comments": "HACKED"})
    job = start_job(client, token(client)).json()
    row = job["extracted_data"]["rows"][0]
    assert row["Amount*"] == "5000"                  # normalised proposal applied
    assert row["Comments"] == MEMBER_DATA["comments"]  # not requested -> untouched
    assert row["_status"] != "ACCEPTED"              # a proposal still needs a human
    log = job["extracted_data"]["repair_log"]
    assert log["requested"] == ["r0_001.Amount*"] and log["applied"][0]["after"] == "5000"


def test_unreadable_repair_value_is_rejected_not_written():
    missing = dict(MEMBER_DATA, amount=None)
    client, app, _, _ = client_for(data={"MEMBER_RECEIPT": missing, "VENDOR_INVOICE": VENDOR_DATA},
                                   verify_results=[{"overall_status": "FAILED", "fields": []}])
    app.state.repair_agent.provider = StaticRepairProvider({"Amount*": "about five thousand"})
    job = start_job(client, token(client)).json()
    assert job["overall_status"] == "NEEDS_REVIEW"
    assert job["extracted_data"]["rows"][0]["Amount*"] == "-"
    assert job["extracted_data"]["repair_log"]["rejected"]


def test_verification_accepts_null_evidence_for_absent_fields():
    from app.accounting.schemas import VerificationResult

    result = VerificationResult(overall_status="PASSED", fields=[
        {"field": "meter_number", "extracted_value": None, "verified_value": None,
         "status": "VERIFIED", "confidence": 1.0, "evidence": None, "page_number": None}])
    assert result.fields[0].evidence == ""
